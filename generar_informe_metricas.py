"""
Informe corto, enfocado solo en métricas y desempeño de los 3 modelos
(Random Forest, XGBoost, LSTM) sobre el dataset real — sin metodología,
objetivos ni contexto del proyecto (para eso ver generar_informe.py).

Uso:
    python generar_informe_metricas.py
"""

import os
import pandas as pd
from docx import Document

from generar_informe import add_title_page, add_metrics_table, add_classification_report_table, add_figure

FIGURES_DIR = "outputs/figures_real"
REPORTS_DIR = "outputs/reports_real"
OUTPUT_PATH = "outputs/reports/informe_metricas_modelos.docx"


def build_report():
    comparison_df = pd.read_csv(os.path.join(REPORTS_DIR, "comparacion_modelos.csv"), index_col=0)
    rf_report = pd.read_csv(os.path.join(REPORTS_DIR, "reporte_clasificacion_random_forest.csv"), index_col=0)
    xgb_report = pd.read_csv(os.path.join(REPORTS_DIR, "reporte_clasificacion_xgboost.csv"), index_col=0)

    doc = Document()
    add_title_page(doc, subtitle="Informe de métricas y desempeño de los modelos")

    doc.add_heading("Cuadro comparativo general", level=1)
    doc.add_paragraph(
        "Random Forest y XGBoost: clasificación baja/media/alta (accuracy, f1_macro). "
        "LSTM: regresión de la constante cinética k (rmse, mae, r2). Corrida sobre "
        "78 curvas reales (nivel de homogeneidad 'núcleo')."
    )
    add_metrics_table(doc, comparison_df)

    doc.add_heading("Random Forest — clasificación", level=1)
    doc.add_paragraph("Precision / recall / f1-score por clase:")
    add_classification_report_table(doc, rf_report)
    doc.add_paragraph()
    add_figure(doc, FIGURES_DIR, "matriz_confusion_random_forest.png", "Matriz de confusión — Random Forest.", width_cm=9)
    add_figure(doc, FIGURES_DIR, "roc_random_forest.png", "Curvas ROC one-vs-rest — Random Forest.", width_cm=9)

    doc.add_heading("XGBoost — clasificación", level=1)
    doc.add_paragraph("Precision / recall / f1-score por clase:")
    add_classification_report_table(doc, xgb_report)
    doc.add_paragraph()
    add_figure(doc, FIGURES_DIR, "matriz_confusion_xgboost.png", "Matriz de confusión — XGBoost.", width_cm=9)
    add_figure(doc, FIGURES_DIR, "roc_xgboost.png", "Curvas ROC one-vs-rest — XGBoost.", width_cm=9)

    doc.add_heading("LSTM — regresión de k", level=1)
    lstm_row = comparison_df.loc["LSTM (regresión de k)"]
    doc.add_paragraph(
        f"RMSE = {lstm_row['rmse']:.5f} día⁻¹  ·  MAE = {lstm_row['mae']:.5f} día⁻¹  ·  R² = {lstm_row['r2']:.3f}"
    )
    add_figure(doc, FIGURES_DIR, "comparacion_curvas_modelos.png",
               "Curva real vs. reconstruida con k predicho (LSTM), y curvas coloreadas por clase predicha (RF/XGBoost).",
               width_cm=17)

    doc.add_heading("Explicabilidad (SHAP) — importancia global", level=1)
    add_figure(doc, FIGURES_DIR, "shap_global_random_forest.png", "Importancia global (SHAP) — Random Forest.", width_cm=13)
    add_figure(doc, FIGURES_DIR, "shap_global_xgboost.png", "Importancia global (SHAP) — XGBoost.", width_cm=13)

    os.makedirs("outputs/reports", exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Informe generado: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_report()
