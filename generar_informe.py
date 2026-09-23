"""
Genera el informe de avance del proyecto en formato Word (.docx), para
presentar a la dirección de tesis. Toma los resultados ya generados por
`main.py` (outputs/reports/comparacion_modelos.csv y outputs/figures/*.png).

Uso:
    python generar_informe.py                # dataset sintético (default)
    python generar_informe.py --source real   # dataset real (HU1)
"""

import argparse
import os
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_title_page(doc, subtitle="Informe de avance"):
    doc.add_paragraph()
    title = doc.add_heading(
        "Aplicación de IA para la predicción de tasas de degradación\n"
        "de polímeros biocompatibles",
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(subtitle)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(16)
    sub.runs[0].font.italic = True

    doc.add_paragraph()
    for line in [
        "Trabajo Final — Carrera de Especialización en Inteligencia Artificial (FIUBA)",
        "Autor: Lic. Gastón Corti",
        "Director: Esp. Ing. Ariadna Garmendia",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_metrics_table(doc, comparison_df):
    table = doc.add_table(rows=1, cols=len(comparison_df.columns) + 1)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Modelo"
    for i, col in enumerate(comparison_df.columns):
        hdr[i + 1].text = col

    for model_name, row in comparison_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(model_name)
        for i, col in enumerate(comparison_df.columns):
            val = row[col]
            cells[i + 1].text = "" if pd.isna(val) else f"{val:.4f}"


def add_classification_report_table(doc, report_df):
    cols = ["precision", "recall", "f1-score", "support"]
    table = doc.add_table(rows=1, cols=len(cols) + 1)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Clase"
    for i, col in enumerate(cols):
        hdr[i + 1].text = col

    for class_name, row in report_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(class_name)
        for i, col in enumerate(cols):
            val = row[col]
            cells[i + 1].text = f"{int(val)}" if col == "support" else f"{val:.3f}"


def add_figure(doc, figures_dir, filename, caption, width_cm=15):
    path = os.path.join(figures_dir, filename)
    if not os.path.exists(path):
        doc.add_paragraph(f"[Figura no encontrada: {filename}]")
        return
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.italic = True
    cap.runs[0].font.size = Pt(10)


# --- Contenido específico por fuente de datos (sintético vs. real) ---

SOURCE_CONTEXT = {
    "sintetico": {
        "subtitle": "Informe de avance",
        "estado_actual": (
            "El pipeline corre de punta a punta (CRISP-DM completo) con datos "
            "sintéticos generados mediante una cinética de pseudo-primer orden "
            "fisicoquímicamente plausible, mientras avanza en paralelo la extracción "
            "real de curvas desde publicaciones científicas (HU1, Épica 1). Esto "
            "permitió desarrollar y validar el resto del pipeline (preprocesamiento, "
            "modelado, evaluación, visualización, tracking de experimentos) sin "
            "depender de tener primero el dataset bibliográfico completo."
        ),
        "datos_desc": (
            "Cada curva incluye: tipo de polímero (PLA, PCL, PGA, PLGA, etc.), medio de "
            "ensayo (SBF, PBS, etc.), temperatura, pH, peso molecular inicial, "
            "cristalinidad, área superficial, y la serie temporal de pérdida de masa."
        ),
        "corrida_desc": "80 curvas sintéticas, semilla 42",
        "resultados_resumen": (
            "Random Forest y XGBoost clasifican con accuracy alto (1.00 y 0.94 "
            "respectivamente) sobre el conjunto de test. El LSTM predice k con un "
            "error absoluto medio (MAE) de ~0.006 día⁻¹ y R²≈0.6, con mejor ajuste "
            "en curvas de degradación rápida que en las de degradación lenta."
        ),
        "shap_resumen": (
            "Ambos modelos coinciden en que el peso molecular inicial (initial_mw_kDa) "
            "es la variable más influyente, seguida de la cristalinidad y el tipo de "
            "polímero (PGA en particular) — coherente con lo esperado fisicoquímicamente: "
            "mayor peso molecular y cristalinidad se asocian a degradación más lenta."
        ),
        "proximos_pasos": [
            "HU1 (Épica 1): extracción real de curvas desde publicaciones científicas "
            "con WebPlotDigitizer, y reemplazo del dataset sintético.",
            "Re-evaluar los 3 modelos sobre el dataset real y comparar contra el "
            "desempeño obtenido con datos sintéticos.",
            "Ajuste de hiperparámetros (GridSearchCV, ya soportado por el pipeline) "
            "sobre el dataset real.",
        ],
    },
    "real": {
        "subtitle": "Informe de avance — dataset real (HU1)",
        "estado_actual": (
            "El pipeline ya corre de punta a punta sobre el dataset REAL extraído de "
            "publicaciones científicas (HU1, Épica 1, completada): 4689 puntos, 120 "
            "curvas, 27 papers, filtrados al subconjunto 'núcleo' (78 curvas) "
            "comparable a medios fisiológicos (PBS/SBF/saliva artificial a ~37°C), "
            "excluyendo condiciones extremas (NaOH, pH extremo) y ensayos acelerados. "
            "El esquema de datos se extendió respecto del sintético para aprovechar "
            "variables no disponibles antes: composición de copolímero, geometría, "
            "método de fabricación y porosidad."
        ),
        "datos_desc": (
            "Cada curva incluye: familia de polímero (incluye blends, ej. PCL/PLLA, "
            "PLLA/PLGA/PCL), fracciones de composición de copolímero (PLLA/PLGA/PCL/PEG, "
            "y lactide:glycolide dentro de PLGA), geometría (scaffold, film, mesh, "
            "fibra, etc.), método de fabricación, porosidad, medio de ensayo, "
            "temperatura, pH, peso molecular inicial, y la serie temporal de pérdida "
            "de masa. No incluye cristalinidad ni área superficial (no reportadas en "
            "la literatura relevada)."
        ),
        "corrida_desc": "78 curvas reales, nivel de homogeneidad 'núcleo'",
        "resultados_resumen": (
            "Random Forest (accuracy 0.75) y XGBoost (accuracy 0.69) clasifican con "
            "desempeño más modesto que sobre datos sintéticos, esperable dado el ruido "
            "y la heterogeneidad real entre 27 estudios distintos. El LSTM predice k "
            "con MAE≈0.002 día⁻¹ y R²≈0.77. Nota: los bins de clasificación "
            "(baja/media/alta) se recalibraron por terciles de la distribución real de "
            "k — los bins originales (pensados para el rango sintético) dejaban la "
            "clase 'alta' vacía, porque el k fisiológico real es 10-100x menor. Esta "
            "recalibración es provisoria y debe validarse con la dirección antes de "
            "reportar resultados finales."
        ),
        "shap_resumen": (
            "A diferencia del dataset sintético, acá emerge una señal fisicoquímica "
            "muy concreta: la fracción de PLGA (plga_fraction) y las fracciones de "
            "láctico/glicólico dentro del PLGA (lactide/glycolide_fraction_in_plga) son "
            "las variables más influyentes en ambos modelos — coincide exactamente con "
            "la literatura: la proporción láctico:glicólico es el factor clásico que "
            "gobierna la velocidad de degradación del PLGA. Esta señal era invisible "
            "con el esquema sintético/mínimo, que solo tenía 'tipo de polímero' como "
            "categoría única."
        ),
        "proximos_pasos": [
            "Validar con la dirección los bins de clasificación baja/media/alta "
            "(actualmente por terciles de k, no por literatura clínica).",
            "Revisar curvas con ajuste cinético de mala calidad (R² bajo o negativo, "
            "ej. por curvas no monótonas o con pocos puntos) y decidir si se filtran.",
            "Evaluar incorporar los niveles 'acelerado' y 'extendido' (con la "
            "normalización a k equivalente a 37°C ya calculada en el dataset) para "
            "ampliar la muestra.",
            "Ajuste de hiperparámetros (GridSearchCV, ya soportado por el pipeline) "
            "sobre el dataset real.",
        ],
    },
}


def build_report(source: str = "sintetico"):
    ctx = SOURCE_CONTEXT[source]

    if source == "real":
        figures_dir = "outputs/figures_real"
        reports_dir = "outputs/reports_real"
        processed_data_path = "outputs/reports_real/dataset_real_snapshot.csv"
    else:
        figures_dir = "outputs/figures"
        reports_dir = "outputs/reports"
        processed_data_path = "data/processed/dataset_estandarizado.csv"

    output_path = os.path.join("outputs/reports", "informe_avance_tesis.docx")

    comparison_df = pd.read_csv(os.path.join(reports_dir, "comparacion_modelos.csv"), index_col=0)
    df_curves = pd.read_csv(processed_data_path)
    class_counts = df_curves["degradation_class"].value_counts()
    rf_report = pd.read_csv(os.path.join(reports_dir, "reporte_clasificacion_random_forest.csv"), index_col=0)
    xgb_report = pd.read_csv(os.path.join(reports_dir, "reporte_clasificacion_xgboost.csv"), index_col=0)

    doc = Document()
    add_title_page(doc, subtitle=ctx["subtitle"])

    doc.add_heading("1. Objetivo del proyecto", level=1)
    doc.add_paragraph(
        "Implementar un pipeline de IA que, a partir de curvas de degradación "
        "(pérdida de masa vs. tiempo) de polímeros biocompatibles en medios "
        "fisiológicos, estime la constante cinética k y clasifique la tasa de "
        "degradación (baja/media/alta), comparando modelos supervisados "
        "(Random Forest, XGBoost) para la clasificación y un modelo de deep "
        "learning (LSTM) para la regresión de k."
    )

    doc.add_heading("2. Estado actual", level=1)
    doc.add_paragraph(ctx["estado_actual"])

    doc.add_heading("3. Metodología (CRISP-DM)", level=1)
    doc.add_paragraph("El pipeline (main.py) ejecuta las siguientes fases en orden:")
    for item in [
        "Carga de datos (sintéticos o reales, según el archivo de configuración).",
        "Estandarización y limpieza de unidades (HU2).",
        "Ajuste cinético de pseudo-primer orden por curva (mass_loss_pct(t) = 100·(1−exp(−k·t))) "
        "y categorización en clases baja/media/alta según el k ajustado.",
        "Visualizaciones exploratorias: distribución de clases, matriz de correlaciones, "
        "muestra de curvas.",
        "Entrenamiento y comparación de modelos, con tracking de parámetros/métricas/artefactos "
        "en MLflow (HU3, HU4, HU5).",
        "Visualización comparativa de las curvas reconstruidas/clasificadas por cada modelo.",
        "Evaluación detallada de clasificación (matrices de confusión, ROC/AUC) y "
        "explicabilidad (SHAP) para Random Forest y XGBoost.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("4. Datos utilizados", level=1)
    doc.add_paragraph(
        f"Dataset actual: {len(df_curves)} curvas, con la siguiente distribución "
        f"de clases de tasa de degradación: baja={class_counts.get('baja', 0)}, "
        f"media={class_counts.get('media', 0)}, alta={class_counts.get('alta', 0)}."
    )
    doc.add_paragraph(ctx["datos_desc"])

    doc.add_heading("5. Modelos implementados", level=1)
    doc.add_paragraph("Clasificación (baja/media/alta), comparando dos modelos supervisados:")
    for item in ["Random Forest (scikit-learn).", "XGBoost."]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Regresión de la constante cinética k, con una red LSTM (PyTorch) que toma la curva "
        "interpolada de pérdida de masa como secuencia de entrada."
    )
    doc.add_paragraph(
        "Todos los entrenamientos quedan registrados en MLflow (parámetros, métricas y el "
        "modelo serializado como artefacto), lo que permite comparar corridas entre sí desde "
        "la interfaz web de MLflow."
    )

    doc.add_heading("6. Resultados", level=1)
    doc.add_paragraph(f"Cuadro comparativo de la corrida más reciente ({ctx['corrida_desc']}):")
    add_metrics_table(doc, comparison_df)
    doc.add_paragraph()
    doc.add_paragraph(ctx["resultados_resumen"])

    doc.add_heading("6.1. Distribución de clases y correlaciones", level=2)
    add_figure(doc, figures_dir, "distribucion_clases.png", "Figura 1. Distribución de clases de tasa de degradación.", width_cm=10)
    add_figure(doc, figures_dir, "matriz_correlaciones.png", "Figura 2. Matriz de correlaciones entre variables numéricas.", width_cm=12)

    doc.add_heading("6.2. Curvas de degradación", level=2)
    add_figure(doc, figures_dir, "curvas_degradacion_muestra.png", "Figura 3. Muestra de curvas de degradación del dataset.", width_cm=14)

    doc.add_heading("6.3. Comparación de modelos sobre curvas de test", level=2)
    doc.add_paragraph(
        "Para una muestra de curvas del conjunto de test: a la izquierda, la curva real "
        "(línea con marcadores) contra la reconstruida con el k predicho por el LSTM (punteada); "
        "al centro y a la derecha, la curva real coloreada según la clase que le asignó cada "
        "clasificador (línea sólida = acierto, punteada = error)."
    )
    add_figure(doc, figures_dir, "comparacion_curvas_modelos.png", "Figura 4. Comparación de predicciones de los 3 modelos sobre curvas de test.", width_cm=17)

    doc.add_heading("6.4. Evaluación detallada de clasificación (RF vs XGBoost)", level=2)
    doc.add_paragraph(
        "Además de accuracy y f1_macro (Sección 6), se detalla la matriz de confusión, las "
        "curvas ROC one-vs-rest con AUC por clase, y el precision/recall/f1-score desglosado "
        "por clase, para ambos clasificadores sobre el mismo conjunto de test."
    )

    doc.add_heading("Random Forest", level=3)
    add_classification_report_table(doc, rf_report)
    doc.add_paragraph()
    add_figure(doc, figures_dir, "matriz_confusion_random_forest.png", "Figura 5. Matriz de confusión — Random Forest.", width_cm=9)
    add_figure(doc, figures_dir, "roc_random_forest.png", "Figura 6. Curvas ROC one-vs-rest — Random Forest.", width_cm=9)

    doc.add_heading("XGBoost", level=3)
    add_classification_report_table(doc, xgb_report)
    doc.add_paragraph()
    add_figure(doc, figures_dir, "matriz_confusion_xgboost.png", "Figura 7. Matriz de confusión — XGBoost.", width_cm=9)
    add_figure(doc, figures_dir, "roc_xgboost.png", "Figura 8. Curvas ROC one-vs-rest — XGBoost.", width_cm=9)

    doc.add_heading("6.5. Explicabilidad (SHAP)", level=2)
    doc.add_paragraph(
        "Se usó SHAP (TreeExplainer, exacto para modelos de árboles) para explicar las "
        "predicciones de Random Forest y XGBoost: importancia global (promedio de |valor SHAP| "
        "entre todas las curvas de test) y una explicación local (por qué el modelo clasificó "
        "una curva particular como lo hizo). " + ctx["shap_resumen"]
    )

    doc.add_heading("Random Forest", level=3)
    add_figure(doc, figures_dir, "shap_global_random_forest.png", "Figura 9. Importancia global (SHAP) — Random Forest.", width_cm=13)
    add_figure(doc, figures_dir, "shap_local_random_forest.png", "Figura 10. Explicación local (SHAP) de una curva de test — Random Forest.", width_cm=15)

    doc.add_heading("XGBoost", level=3)
    add_figure(doc, figures_dir, "shap_global_xgboost.png", "Figura 11. Importancia global (SHAP) — XGBoost.", width_cm=13)
    add_figure(doc, figures_dir, "shap_local_xgboost.png", "Figura 12. Explicación local (SHAP) de una curva de test — XGBoost.", width_cm=15)

    doc.add_heading("7. Próximos pasos", level=1)
    for item in ctx["proximos_pasos"]:
        doc.add_paragraph(item, style="List Bullet")

    os.makedirs("outputs/reports", exist_ok=True)
    doc.save(output_path)
    print(f"Informe generado: {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", choices=["sintetico", "real"], default="sintetico")
    args = parser.parse_args()
    build_report(args.source)
