"""
Informe comparativo: desempeño de los modelos sobre el dataset SINTÉTICO vs.
el dataset REAL, lado a lado. Requiere que existan resultados de ambas
corridas: outputs/reports/ + outputs/figures/ (sintético, la corrida más
reciente) y outputs/reports_real/ + outputs/figures_real/ (real, respaldados
aparte porque ambas corridas escriben a las mismas rutas por defecto).

Uso:
    python generar_informe_comparativo.py
"""

import os
import pandas as pd
from docx import Document

from generar_informe import add_title_page, add_figure

OUTPUT_PATH = "outputs/reports/informe_comparativo_sintetico_vs_real.docx"


def add_side_by_side_table(doc, comp_synth, comp_real):
    metrics = ["accuracy", "f1_macro", "rmse", "mae", "r2"]
    models = list(dict.fromkeys(list(comp_synth.index) + list(comp_real.index)))

    table = doc.add_table(rows=1, cols=2 + len(metrics))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Modelo"
    hdr[1].text = "Dataset"
    for i, m in enumerate(metrics):
        hdr[2 + i].text = m

    for model in models:
        for label, df in [("Sintético", comp_synth), ("Real", comp_real)]:
            if model not in df.index:
                continue
            row = df.loc[model]
            cells = table.add_row().cells
            cells[0].text = model
            cells[1].text = label
            for i, m in enumerate(metrics):
                val = row.get(m)
                cells[2 + i].text = "" if pd.isna(val) else f"{val:.4f}"


def build_report():
    comp_synth = pd.read_csv("outputs/reports/comparacion_modelos.csv", index_col=0)
    comp_real = pd.read_csv("outputs/reports_real/comparacion_modelos.csv", index_col=0)

    doc = Document()
    add_title_page(doc, subtitle="Comparación de desempeño: dataset sintético vs. real")

    doc.add_heading("1. Cuadro comparativo", level=1)
    add_side_by_side_table(doc, comp_synth, comp_real)

    doc.add_paragraph()
    doc.add_paragraph(
        "Nota sobre RMSE/MAE del LSTM: no son directamente comparables entre "
        "datasets porque la escala de k difiere mucho (k sintético llega a ser "
        "varias veces mayor que el k real fisiológico) — un RMSE más chico en el "
        "real no implica necesariamente un mejor ajuste relativo. El R² sí es "
        "comparable al ser una medida relativa al propio rango de cada dataset."
    )

    doc.add_heading("2. Lectura de los resultados", level=1)
    doc.add_paragraph(
        "Random Forest y XGBoost pierden accuracy al pasar de datos sintéticos "
        "(1.00 / 0.94) a datos reales (0.75 / 0.69). Es el resultado esperado: el "
        "dataset sintético fue generado con el mismo modelo funcional "
        "(pseudo-primer orden) que después se usa para ajustar k y clasificar, sin "
        "ruido entre estudios; el dataset real combina 27 papers con protocolos, "
        "instrumentos y calidad de digitalización distintos, además de menos curvas "
        "(78 vs. 80, pero mucho más heterogéneas)."
    )
    doc.add_paragraph(
        "El LSTM, en cambio, mejora su R² (0.57 → 0.77) sobre datos reales. Una "
        "lectura posible es que las curvas reales del subconjunto 'núcleo' son, en "
        "su mayoría, más largas y densamente muestreadas (hasta 100 puntos) que las "
        "sintéticas (6-14 puntos), lo que le da a la red más señal temporal por "
        "curva para aprender la forma de la curva, compensando el ruido entre "
        "estudios."
    )

    doc.add_heading("3. Explicabilidad: la diferencia más importante", level=1)
    doc.add_paragraph(
        "Más allá de las métricas, la diferencia cualitativa más relevante está en "
        "SHAP: con el esquema sintético (polymer_type como única variable de "
        "composición), la variable más influyente es el peso molecular inicial. Con "
        "el esquema real extendido (fracciones de copolímero, geometría, método de "
        "fabricación, porosidad), emerge una señal mucho más específica: la "
        "fracción de PLGA y la proporción láctico:glicólico dentro del PLGA — "
        "exactamente el factor que la literatura identifica como determinante de la "
        "velocidad de degradación de ese copolímero. Esto valida la decisión de "
        "extender el esquema en vez de mapear los datos reales al esquema mínimo "
        "sintético: esa señal hubiera quedado invisible."
    )

    doc.add_heading("3.1. Importancia global (SHAP) — Random Forest", level=2)
    add_figure(doc, "outputs/figures", "shap_global_random_forest.png", "Sintético: domina peso molecular y cristalinidad.", width_cm=13)
    add_figure(doc, "outputs/figures_real", "shap_global_random_forest.png", "Real: domina la composición de copolímero (PLGA/láctico/glicólico).", width_cm=13)

    doc.add_heading("4. Conclusión", level=1)
    doc.add_paragraph(
        "El pipeline es robusto a ambas fuentes de datos sin cambios de código "
        "(mismo esquema parametrizable, mismos modelos, misma evaluación). El salto "
        "a datos reales bajó el accuracy de clasificación, como es esperable, pero "
        "generó hallazgos fisicoquímicamente más ricos y creíbles vía SHAP. "
        "Próximo paso natural: validar con la dirección los bins de clasificación y "
        "considerar ampliar la muestra con los niveles 'acelerado'/'extendido' del "
        "dataset real."
    )

    os.makedirs("outputs/reports", exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Informe generado: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_report()
